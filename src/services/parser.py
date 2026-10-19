from pathlib import Path 
import docx
from pypdf import PdfReader

def is_supported_file(filename: str) -> bool:
    suffix = Path(filename).suffix.lower()
    return suffix in (".txt", ".docx", ".pdf")

def extract_text_from_file(file_path: Path) -> str:
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()

    if suffix == ".txt":
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read().strip()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="cp1251", errors="ignore") as f:
                return f.read().strip()
    
    elif suffix == ".docx":
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text).strip()
    
    elif suffix == ".pdf":
        reader = PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            text = page.extract_text() 
            if text:
                pages_text.append(text.strip())
        return "\n".join(pages_text).strip()
    else:
        raise ValueError(f"Непідтримуваний формат файлу: {suffix}")                   
